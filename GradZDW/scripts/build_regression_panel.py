#!/usr/bin/env python3
"""Build the first regression panel for the thesis."""

from __future__ import annotations

import csv
import math
from pathlib import Path

import openpyxl


BANK_DATA = Path("02_数据集/上市银行年报PDF/上市银行年报指标提取_v4_保守补流动性比例.csv")
DFI_DATA = Path("02_数据集/北京大学数字普惠金融指数（PKU-DFIIC）2011-2023.xlsx")
OUT_CSV = Path("02_数据集/回归面板数据_流动性风险_2014-2023_初版.csv")
OUT_XLSX = Path("02_数据集/回归面板数据_流动性风险_2014-2023_初版.xlsx")
MAPPING_CSV = Path("02_数据集/上市银行总部地区匹配表.csv")


BANK_META = {
    "000001": {"bank_name": "平安银行", "hq_city": "深圳市", "hq_province": "广东省", "bank_type": "股份制商业银行"},
    "600000": {"bank_name": "浦发银行", "hq_city": "上海市", "hq_province": "上海市", "bank_type": "股份制商业银行"},
    "600016": {"bank_name": "民生银行", "hq_city": "北京市", "hq_province": "北京市", "bank_type": "股份制商业银行"},
    "600036": {"bank_name": "招商银行", "hq_city": "深圳市", "hq_province": "广东省", "bank_type": "股份制商业银行"},
    "601166": {"bank_name": "兴业银行", "hq_city": "福州市", "hq_province": "福建省", "bank_type": "股份制商业银行"},
    "601288": {"bank_name": "农业银行", "hq_city": "北京市", "hq_province": "北京市", "bank_type": "国有大型商业银行"},
    "601328": {"bank_name": "交通银行", "hq_city": "上海市", "hq_province": "上海市", "bank_type": "国有大型商业银行"},
    "601398": {"bank_name": "工商银行", "hq_city": "北京市", "hq_province": "北京市", "bank_type": "国有大型商业银行"},
    "601939": {"bank_name": "建设银行", "hq_city": "北京市", "hq_province": "北京市", "bank_type": "国有大型商业银行"},
    "601988": {"bank_name": "中国银行", "hq_city": "北京市", "hq_province": "北京市", "bank_type": "国有大型商业银行"},
}


def to_float(value: str) -> float | None:
    value = (value or "").strip().replace(",", "")
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def load_city_dfi() -> dict[tuple[int, str], dict[str, float | None]]:
    wb = openpyxl.load_workbook(DFI_DATA, read_only=True, data_only=True)
    ws = wb["Prefecture_Level_Cities"]
    header = list(next(ws.iter_rows(values_only=True)))
    data: dict[tuple[int, str], dict[str, float | None]] = {}
    for row in ws.iter_rows(values_only=True):
        rec = dict(zip(header, row))
        year = rec.get("year")
        city = rec.get("pref_name_year18")
        if not isinstance(year, int) or not city:
            continue
        data[(year, city)] = {
            "dfi_index": rec.get("index_aggregate"),
            "dfi_breadth": rec.get("coverage_breadth"),
            "dfi_depth": rec.get("usage_depth"),
            "dfi_digitization": rec.get("digitization_level"),
            "dfi_payment": rec.get("payment"),
            "dfi_credit": rec.get("credit"),
            "dfi_investment": rec.get("investment"),
            "dfi_insurance": rec.get("insurance"),
        }
    return data


def save_mapping() -> None:
    with MAPPING_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        fields = ["code", "bank_name", "hq_city", "hq_province", "bank_type"]
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for code, meta in BANK_META.items():
            writer.writerow({"code": code, **meta})


def main() -> int:
    save_mapping()
    city_dfi = load_city_dfi()
    rows = list(csv.DictReader(BANK_DATA.open(encoding="utf-8-sig")))
    out_rows = []
    for row in rows:
        code = row["code"]
        year = int(row["year"])
        meta = BANK_META.get(code)
        if not meta:
            continue
        dfi = city_dfi.get((year, meta["hq_city"]))
        if not dfi:
            continue

        liquidity_ratio = to_float(row.get("liquidity_ratio", ""))
        total_assets = to_float(row.get("total_assets", ""))
        roa = to_float(row.get("roa", ""))
        roe = to_float(row.get("roe", ""))
        car = to_float(row.get("capital_adequacy_ratio", ""))
        core_car = to_float(row.get("core_tier1_capital_ratio", ""))
        npl = to_float(row.get("npl_ratio", ""))
        cost_income = to_float(row.get("cost_income_ratio", ""))
        provision = to_float(row.get("provision_coverage", ""))

        out = {
            "code": code,
            "bank_name": row["bank_name"],
            "year": year,
            "hq_city": meta["hq_city"],
            "hq_province": meta["hq_province"],
            "bank_type": meta["bank_type"],
            "is_state_owned_large": 1 if meta["bank_type"] == "国有大型商业银行" else 0,
            "liquidity_ratio": liquidity_ratio,
            "liquidity_risk": -liquidity_ratio if liquidity_ratio is not None else None,
            "total_assets": total_assets,
            "ln_total_assets": math.log(total_assets) if total_assets and total_assets > 0 else None,
            "roa": roa,
            "roe": roe,
            "capital_adequacy_ratio": car,
            "core_tier1_capital_ratio": core_car,
            "npl_ratio": npl,
            "cost_income_ratio": cost_income,
            "provision_coverage": provision,
            **dfi,
            "liquidity_ratio_page": row.get("liquidity_ratio_page", ""),
            "liquidity_ratio_evidence": row.get("liquidity_ratio_evidence", ""),
            "pdf_file": row.get("pdf_file", ""),
        }
        out_rows.append(out)

    fields = [
        "code",
        "bank_name",
        "year",
        "hq_city",
        "hq_province",
        "bank_type",
        "is_state_owned_large",
        "liquidity_ratio",
        "liquidity_risk",
        "dfi_index",
        "dfi_breadth",
        "dfi_depth",
        "dfi_digitization",
        "dfi_payment",
        "dfi_credit",
        "dfi_investment",
        "dfi_insurance",
        "total_assets",
        "ln_total_assets",
        "roa",
        "roe",
        "capital_adequacy_ratio",
        "core_tier1_capital_ratio",
        "npl_ratio",
        "cost_income_ratio",
        "provision_coverage",
        "liquidity_ratio_page",
        "liquidity_ratio_evidence",
        "pdf_file",
    ]
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(out_rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "panel"
    ws.append(fields)
    for row in out_rows:
        ws.append([row.get(field) for field in fields])
    for col in ws.columns:
        max_len = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 42)
        ws.column_dimensions[col[0].column_letter].width = max_len
    wb.save(OUT_XLSX)

    print(f"Rows: {len(out_rows)}")
    print(f"Saved: {OUT_CSV}")
    print(f"Saved: {OUT_XLSX}")
    print(f"Saved: {MAPPING_CSV}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
